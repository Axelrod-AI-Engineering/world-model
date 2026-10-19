"""Parsing (Oct 17): a file's bytes into plain text, deterministically.

Each format has one parser, chosen by the file's extension (then its MIME type). Text is normalized the same way for
every format: Unicode NFC, curly quotes and dashes kept, whitespace runs inside a line collapsed, blank-line runs
collapsed to one. Pages, sheets, and slides become `segments`: a document's text is their texts joined by blank
lines, and each segment's span is kept so a chunk knows its page.

Formats: PDF (PDFium, through pypdfium2: fast, and it reads PDFs whose fonts pypdf could not), DOCX and PPTX (their XML, read from the zip), XLSX (openpyxl, values not formulas), HTML
(the standard library's parser), CSV, Markdown and plain text, and JSON (an import bundle is passed through whole).
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePath

PARSER_VERSION = "parse-1.0"


class ParseError(ValueError):
    """The file could not be read as its format says it is."""


@dataclass(frozen=True)
class Segment:
    text: str
    page: int | None = None          # 1-based page, sheet, or slide number
    label: str = ""                  # e.g. a sheet's name


@dataclass(frozen=True)
class Parsed:
    format: str
    segments: tuple[Segment, ...]
    title: str = ""
    meta: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.segments if s.text)

    def spans(self) -> list[tuple[int, int, Segment]]:
        """Each non-empty segment's character span in `text`."""
        out, at = [], 0
        for s in self.segments:
            if not s.text:
                continue
            out.append((at, at + len(s.text), s))
            at += len(s.text) + 2
        return out


_SPACES = re.compile(r"[ \t  -​  　]+")
_BLANKS = re.compile(r"\n{3,}")


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    lines = [_SPACES.sub(" ", ln).strip() for ln in text.split("\n")]
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


# -- formats ----------------------------------------------------------------------------------------------------
def _pdf(data: bytes) -> Parsed:
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as err:
        raise ParseError(f"not a readable PDF: {err}") from err
    try:
        pages = [normalize(doc[k].get_textpage().get_text_range()) for k in range(len(doc))]
        title = normalize(str(doc.get_metadata_dict().get("Title") or ""))
    finally:
        doc.close()
    return Parsed("pdf", tuple(Segment(t, page=k + 1) for k, t in enumerate(pages)), title=title,
                  meta={"pages": len(pages)})


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _zip(data: bytes) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as err:
        raise ParseError(f"not a valid Office file: {err}") from err


def _docx(data: bytes) -> Parsed:
    from xml.etree import ElementTree as ET

    z = _zip(data)
    try:
        root = ET.fromstring(z.read("word/document.xml"))
    except (KeyError, ET.ParseError) as err:
        raise ParseError(f"not a valid DOCX: {err}") from err
    paras = []
    for p in root.iter(f"{_W}p"):
        runs = []
        for node in p.iter():
            if node.tag == f"{_W}t" and node.text:
                runs.append(node.text)
            elif node.tag in (f"{_W}tab",):
                runs.append(" ")
            elif node.tag in (f"{_W}br", f"{_W}cr"):
                runs.append("\n")
        paras.append("".join(runs))
    text = normalize("\n".join(paras))
    first = next((ln for ln in text.split("\n") if ln.strip()), "")
    return Parsed("docx", (Segment(text),), title=first[:160])


def _pptx(data: bytes) -> Parsed:
    from xml.etree import ElementTree as ET

    z = _zip(data)
    names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                   key=lambda n: int(re.search(r"(\d+)", n).group(1)))
    if not names:
        raise ParseError("not a valid PPTX: no slides")
    segs = []
    for k, n in enumerate(names):
        try:
            root = ET.fromstring(z.read(n))
        except ET.ParseError as err:
            raise ParseError(f"not a valid PPTX: {n}: {err}") from err
        paras = ["".join(t.text or "" for t in p.iter(f"{_A}t")) for p in root.iter(f"{_A}p")]
        segs.append(Segment(normalize("\n".join(paras)), page=k + 1))
    first = next((s.text.split("\n")[0] for s in segs if s.text), "")
    return Parsed("pptx", tuple(segs), title=first[:160], meta={"slides": len(segs)})


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        r = repr(round(v, 10))
        return r if "e" in r else r.rstrip("0").rstrip(".")      # 3e-10 stays 3e-10 (stripping made it 3e-1)
    return " ".join(str(v).split("\n")).replace("\r", " ")         # a wrapped cell stays on its row


def _xlsx(data: bytes) -> Parsed:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as err:
        raise ParseError(f"not a valid XLSX: {err}") from err
    segs = []
    for k, ws in enumerate(wb.worksheets):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(" | ".join(cells))
        segs.append(Segment(normalize(f"Sheet: {ws.title}\n" + "\n".join(rows)) if rows else "", page=k + 1,
                            label=ws.title))
    wb.close()
    return Parsed("xlsx", tuple(segs), meta={"sheets": [s.label for s in segs]})


class _Text(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table",
             "ul", "ol", "header", "footer", "blockquote", "pre"}      # td/th: " | " within their row's line
    SKIP = {"script", "style", "noscript", "template", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.title: list[str] = []
        self.in_title = False

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag == "title":
            self.in_title = True
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")
        if tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self.in_title = False
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title.append(data)
        elif not self.skip:
            self.out.append(data)


def _html(data: bytes) -> Parsed:
    p = _Text()
    p.feed(_decode(data))
    p.close()
    text = normalize(re.sub(r"(\n *\| *)", "\n", "".join(p.out)))
    return Parsed("html", (Segment(text),), title=normalize("".join(p.title))[:160])


def _decode(data: bytes) -> str:
    """UTF-16 when the bytes say so (a byte-order mark), else UTF-8, else Latin-1 (which never fails)."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _csv(data: bytes) -> Parsed:
    rows = [" | ".join(r).strip() for r in csv.reader(io.StringIO(_decode(data)))]
    return Parsed("csv", (Segment(normalize("\n".join(r for r in rows if r.strip("| ")))),))


def _text(data: bytes, fmt: str = "text") -> Parsed:
    text = normalize(_decode(data))
    first = next((ln.lstrip("# ").strip() for ln in text.split("\n") if ln.strip()), "")
    return Parsed(fmt, (Segment(text),), title=first[:160])


def _json(data: bytes) -> Parsed:
    try:
        obj = json.loads(_decode(data))
    except json.JSONDecodeError as err:
        raise ParseError(f"not valid JSON: {err}") from err
    if isinstance(obj, dict) and obj.get("format") == "axelrod.import/v1":
        items = obj.get("items") or []
        text = "\n\n".join(normalize(f"{i.get('headline', '')}\n{i.get('summary', '')}") for i in items if isinstance(i, dict))
        return Parsed("bundle", (Segment(text),), title=str(obj.get("title", ""))[:160], meta={"bundle": obj})
    return Parsed("json", (Segment(normalize(json.dumps(obj, indent=1, sort_keys=True, ensure_ascii=False))),))


PARSERS = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".html": _html, ".htm": _html,
           ".csv": _csv, ".md": lambda d: _text(d, "markdown"), ".markdown": lambda d: _text(d, "markdown"),
           ".txt": _text, ".text": _text, ".json": _json}
MIMES = {"application/pdf": ".pdf", "text/html": ".html", "text/csv": ".csv", "text/markdown": ".md",
         "text/plain": ".txt", "application/json": ".json",
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
         "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx"}


def mime_of(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    return next((m for m, e in MIMES.items() if e == ext), "application/octet-stream")


def supported(filename: str, mime: str = "") -> bool:
    return PurePath(filename).suffix.lower() in PARSERS or mime in MIMES


def parse(data: bytes, filename: str, mime: str = "") -> Parsed:
    """The document's text. Raises ParseError for an unsupported format or a file that is not what it claims."""
    ext = PurePath(filename).suffix.lower()
    fn = PARSERS.get(ext) or PARSERS.get(MIMES.get(mime, ""))
    if fn is None:
        raise ParseError(f"unsupported file type: {ext or mime or 'unknown'}")
    return fn(data)
