"""Parsing (Oct 17): a file's bytes into plain text, deterministically.

Each format has one parser, chosen by the file's extension (then its MIME type). Text is normalized the same way for
every format: Unicode NFC, curly quotes and dashes kept, whitespace runs inside a line collapsed, blank-line runs
collapsed to one. Pages, sheets, and slides become `segments`: a document's text is their texts joined by blank
lines, and each segment's span is kept so a chunk knows its page.

Formats: PDF (PDFium, through pypdfium2: fast, and it reads PDFs whose fonts pypdf could not), DOCX and PPTX (their XML, read from the zip), XLSX (openpyxl, values not formulas), HTML
(the standard library's parser), CSV, Markdown and plain text, and JSON (an import bundle is passed through whole).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePath

from axelrod_ingest.cjk import width
from axelrod_ingest.encoding import Decoded, decode
from axelrod_ingest.ocr import DPI, Ocr, OcrError, pgm_of

PARSER_VERSION = "parse-1.1"


class ParseError(ValueError):
    """The file could not be read as its format says it is."""


@dataclass(frozen=True)
class Segment:
    text: str
    page: int | None = None          # 1-based page, sheet, or slide number
    label: str = ""                  # e.g. a sheet's name


@dataclass(frozen=True)
class Parsed:
    format: str
    segments: tuple[Segment, ...]
    title: str = ""
    meta: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.segments if s.text)

    def spans(self) -> list[tuple[int, int, Segment]]:
        """Each non-empty segment's character span in `text`."""
        out, at = [], 0
        for s in self.segments:
            if not s.text:
                continue
            out.append((at, at + len(s.text), s))
            at += len(s.text) + 2
        return out


_SPACES = re.compile(r"[ \t  -​  　]+")
_BLANKS = re.compile(r"\n{3,}")


def normalize(text: str) -> str:
    text = width(unicodedata.normalize("NFC", text)).replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    lines = [_SPACES.sub(" ", ln).strip() for ln in text.split("\n")]
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


# -- formats ----------------------------------------------------------------------------------------------------
MIN_TEXT = 20                    # fewer letters and digits than this on a page that is a picture: it has no text layer
IMAGE_SHARE = 0.25               # a page is a picture when one image covers at least this share of it


def _is_picture(page) -> bool:
    """Whether a page is a scan: a large image on it (the text layer, if any, is a footer or a watermark)."""
    import pypdfium2.raw as pdfium_c

    width, height = page.get_size()
    return any((r - left) * (top - b) >= IMAGE_SHARE * width * height for left, b, r, top in (
        obj.get_bounds() for obj in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE])))


def _pgm(page) -> bytes:
    """The page drawn as a grayscale PGM at the OCR's resolution."""
    bitmap = page.render(scale=DPI / 72, grayscale=True)
    return pgm_of(bitmap.width, bitmap.height, bitmap.stride, bytes(bitmap.buffer))


def _pdf(data: bytes, ocr: Ocr | None = None) -> Parsed:
    """PDF, page by page. A page that is a picture with no text layer is read by `ocr` when there is one; otherwise (or
    where the engine fails) its number goes in `meta["needs_ocr"]`, so a scan never parses to nothing without a word."""
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as err:
        raise ParseError(f"not a readable PDF: {err}") from err
    sha = hashlib.sha256(data).hexdigest()
    pages: list[str] = []
    needs: list[int] = []
    done: list[int] = []
    try:
        for k in range(len(doc)):
            page = doc[k]
            text = normalize(page.get_textpage().get_text_range())
            if sum(c.isalnum() for c in text) < MIN_TEXT and _is_picture(page):
                try:
                    if ocr is None:
                        raise OcrError("no OCR engine")
                    text = normalize(ocr.page_text(sha, k + 1, lambda page=page: _pgm(page)))
                    done.append(k + 1)
                except OcrError:
                    needs.append(k + 1)
            pages.append(text)
        info = doc.get_metadata_dict()
        title = normalize(str(info.get("Title") or ""))
        author = normalize(str(info.get("Author") or ""))
    finally:
        doc.close()
    meta: dict = {"pages": len(pages)}
    if done:
        meta["ocr"] = {**(ocr.meta if ocr else {}), "pages": done}
    if needs:
        meta["needs_ocr"] = needs
    if author:
        meta["publisher"] = author[:160]
    return Parsed("pdf", tuple(Segment(t, page=k + 1) for k, t in enumerate(pages)), title=title, meta=meta)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _zip(data: bytes) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as err:
        raise ParseError(f"not a valid Office file: {err}") from err


def _docx(data: bytes) -> Parsed:
    from xml.etree import ElementTree as ET

    z = _zip(data)
    try:
        root = ET.fromstring(z.read("word/document.xml"))
    except (KeyError, ET.ParseError) as err:
        raise ParseError(f"not a valid DOCX: {err}") from err
    paras = []
    for p in root.iter(f"{_W}p"):
        runs = []
        for node in p.iter():
            if node.tag == f"{_W}t" and node.text:
                runs.append(node.text)
            elif node.tag in (f"{_W}tab",):
                runs.append(" ")
            elif node.tag in (f"{_W}br", f"{_W}cr"):
                runs.append("\n")
        paras.append("".join(runs))
    text = normalize("\n".join(paras))
    first = next((ln for ln in text.split("\n") if ln.strip()), "")
    return Parsed("docx", (Segment(text),), title=first[:160])


def _pptx(data: bytes) -> Parsed:
    from xml.etree import ElementTree as ET

    z = _zip(data)
    names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                   key=lambda n: int(re.search(r"(\d+)", n).group(1)))
    if not names:
        raise ParseError("not a valid PPTX: no slides")
    segs = []
    for k, n in enumerate(names):
        try:
            root = ET.fromstring(z.read(n))
        except ET.ParseError as err:
            raise ParseError(f"not a valid PPTX: {n}: {err}") from err
        paras = ["".join(t.text or "" for t in p.iter(f"{_A}t")) for p in root.iter(f"{_A}p")]
        segs.append(Segment(normalize("\n".join(paras)), page=k + 1))
    first = next((s.text.split("\n")[0] for s in segs if s.text), "")
    return Parsed("pptx", tuple(segs), title=first[:160], meta={"slides": len(segs)})


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        r = repr(round(v, 10))
        return r if "e" in r else r.rstrip("0").rstrip(".")      # 3e-10 stays 3e-10 (stripping made it 3e-1)
    return " ".join(str(v).split("\n")).replace("\r", " ")         # a wrapped cell stays on its row


def _xlsx(data: bytes) -> Parsed:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as err:
        raise ParseError(f"not a valid XLSX: {err}") from err
    segs = []
    for k, ws in enumerate(wb.worksheets):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(" | ".join(cells))
        segs.append(Segment(normalize(f"Sheet: {ws.title}\n" + "\n".join(rows)) if rows else "", page=k + 1,
                            label=ws.title))
    wb.close()
    return Parsed("xlsx", tuple(segs), meta={"sheets": [s.label for s in segs]})


def _encoding(d: Decoded) -> dict[str, str]:
    """How the bytes were read, for the document's metadata: a fallback to Latin-1 is visible there, never silent."""
    return {"name": d.encoding, "how": d.how}


@dataclass
class _Table:
    """A table being read: finished rows, the row and the cell open now, and the cells a rowspan carries into later rows."""

    rows: list[list[str]] = field(default_factory=list)
    row: list[str] | None = None
    cell: list[str] | None = None
    span: tuple[int, int] = (1, 1)                       # the open cell's colspan and rowspan
    pending: dict[int, list] = field(default_factory=dict)      # column -> [text, rows still to cover]

    def fill(self) -> None:
        """Put the cells a rowspan carries down into this row, from where it stands."""
        assert self.row is not None
        while (col := len(self.row)) in self.pending:
            text, left = self.pending.pop(col)
            self.row.append(text)
            if left > 1:
                self.pending[col] = [text, left - 1]


MAX_SPAN = 30                    # a colspan or rowspan beyond this is a typo
PUBLISHER_META = ("publisher", "og:site_name", "dc.publisher", "application-name", "author")   # a page's <meta> words, best first
DATA_CELL_CHARS = 160           # a table of cells longer than this on average lays out a page; it is not data


def _span(attrs: list, name: str) -> int:
    value = dict(attrs).get(name) or ""
    return min(int(value), MAX_SPAN) if value.isdigit() and int(value) > 0 else 1


class _Text(HTMLParser):
    """HTML to text. A data table becomes one line per row ("Sales Revenue | Current Period | 49,215,328 | 45,938,886"): a cell's
    own line breaks become spaces, a colspan leaves empty cells, a rowspan's text is carried into the rows it covers, so each
    row's label and figures stay together. A table of long cells (a page laid out in a table) is read as blocks, as text."""

    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table",
             "ul", "ol", "header", "footer", "blockquote", "pre"}
    SKIP = {"script", "style", "noscript", "template", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.title: list[str] = []
        self.in_title = False
        self.tables: list[_Table] = []
        self.publisher: dict[str, str] = {}              # the page's own word for who published it (a <meta> tag)

    def _sink(self) -> list[str]:
        """Where text goes: the open cell of the innermost table, else the page."""
        cell = self.tables[-1].cell if self.tables else None
        return self.out if cell is None else cell

    def _break(self) -> str:
        return "\n" if not self.tables or self.tables[-1].cell is None else " "

    def _open_row(self) -> None:
        t = self.tables[-1]
        if t.cell is not None:
            self._close_cell()
        if t.row is not None:
            self._close_row()
        t.row = []
        t.fill()

    def _open_cell(self, attrs: list) -> None:
        t = self.tables[-1]
        if t.cell is not None:
            self._close_cell()
        if t.row is None:
            self._open_row()
        t.cell, t.span = [], (_span(attrs, "colspan"), _span(attrs, "rowspan"))

    def _close_cell(self) -> None:
        t = self.tables[-1]
        assert t.cell is not None and t.row is not None
        text = " ".join("".join(t.cell).split())
        t.fill()
        col, (colspan, rowspan) = len(t.row), t.span
        t.row += [text] + [""] * (colspan - 1)
        if rowspan > 1:
            t.pending.update({col + j: [text if j == 0 else "", rowspan - 1] for j in range(colspan)})
        t.cell = None

    def _close_row(self) -> None:
        t = self.tables[-1]
        assert t.row is not None
        t.fill()
        if any(t.row):
            t.rows.append(t.row)
        t.row = None

    def _close_table(self) -> None:
        t = self.tables[-1]
        if t.cell is not None:
            self._close_cell()
        if t.row is not None:
            self._close_row()
        self.tables.pop()
        cells = [c for row in t.rows for c in row if c]
        wide = any(sum(bool(c) for c in row) >= 2 for row in t.rows)
        data = wide and sum(map(len, cells)) <= DATA_CELL_CHARS * len(cells)
        lines = [" | ".join(row).rstrip(" |") for row in t.rows] if data else cells
        self._sink().append(self._break() + self._break().join(lines) + self._break())

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag == "meta":
            a = {k.lower(): (v or "").strip() for k, v in attrs}
            name = (a.get("name") or a.get("property") or "").lower()
            if name in PUBLISHER_META and a.get("content"):
                self.publisher.setdefault(name, a["content"])
        if tag == "title":
            self.in_title = True
        if tag in self.SKIP:
            self.skip += 1
        elif tag == "table":
            self.tables.append(_Table())
        elif tag == "tr" and self.tables:
            self._open_row()
        elif tag in ("td", "th") and self.tables:
            self._open_cell(attrs)
        elif tag in self.BLOCK:
            self._sink().append(self._break())

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self.in_title = False
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in ("td", "th") and self.tables and self.tables[-1].cell is not None:
            self._close_cell()
        elif tag == "tr" and self.tables and self.tables[-1].row is not None:
            if self.tables[-1].cell is not None:
                self._close_cell()
            self._close_row()
        elif tag == "table" and self.tables:
            self._close_table()
        elif tag in self.BLOCK and not (self.tables and tag in ("table", "tr")):
            self._sink().append(self._break())

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title.append(data)
        elif not self.skip:
            self._sink().append(data)

    def close(self) -> None:
        super().close()
        while self.tables:                               # a table the page never closed
            self._close_table()


def _html(data: bytes) -> Parsed:
    p = _Text()
    d = decode(data, html=True)
    p.feed(d.text)
    p.close()
    text = normalize("".join(p.out))
    meta = {"encoding": _encoding(d)}
    if publisher := next((p.publisher[k] for k in PUBLISHER_META if k in p.publisher), ""):
        meta["publisher"] = normalize(publisher)[:160]
    return Parsed("html", (Segment(text),), title=normalize("".join(p.title))[:160], meta=meta)


def _csv(data: bytes) -> Parsed:
    d = decode(data)
    rows = [" | ".join(r).strip() for r in csv.reader(io.StringIO(d.text))]
    return Parsed("csv", (Segment(normalize("\n".join(r for r in rows if r.strip("| ")))),), meta={"encoding": _encoding(d)})


def _text(data: bytes, fmt: str = "text") -> Parsed:
    d = decode(data)
    text = normalize(d.text)
    first = next((ln.lstrip("# ").strip() for ln in text.split("\n") if ln.strip()), "")
    return Parsed(fmt, (Segment(text),), title=first[:160], meta={"encoding": _encoding(d)})


def _json(data: bytes) -> Parsed:
    d = decode(data)
    try:
        obj = json.loads(d.text)
    except json.JSONDecodeError as err:
        raise ParseError(f"not valid JSON: {err}") from err
    if isinstance(obj, dict) and obj.get("format") == "axelrod.import/v1":
        items = obj.get("items") or []
        text = "\n\n".join(normalize(f"{i.get('headline', '')}\n{i.get('summary', '')}") for i in items if isinstance(i, dict))
        return Parsed("bundle", (Segment(text),), title=str(obj.get("title", ""))[:160],
                      meta={"bundle": obj, "encoding": _encoding(d)})
    return Parsed("json", (Segment(normalize(json.dumps(obj, indent=1, sort_keys=True, ensure_ascii=False))),),
                  meta={"encoding": _encoding(d)})


PARSERS = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".html": _html, ".htm": _html,
           ".csv": _csv, ".md": lambda d: _text(d, "markdown"), ".markdown": lambda d: _text(d, "markdown"),
           ".txt": _text, ".text": _text, ".json": _json}
MIMES = {"application/pdf": ".pdf", "text/html": ".html", "text/csv": ".csv", "text/markdown": ".md",
         "text/plain": ".txt", "application/json": ".json",
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
         "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx"}


def mime_of(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    return next((m for m, e in MIMES.items() if e == ext), "application/octet-stream")


def supported(filename: str, mime: str = "") -> bool:
    return PurePath(filename).suffix.lower() in PARSERS or mime in MIMES


def parse(data: bytes, filename: str, mime: str = "", *, ocr: Ocr | None = None) -> Parsed:
    """The document's text. Raises ParseError for an unsupported format or a file that is not what it claims. `ocr`
    reads the pages of a PDF that are pictures (see ocr.py); without it they are listed in `meta["needs_ocr"]`."""
    ext = PurePath(filename).suffix.lower()
    fn = PARSERS.get(ext) or PARSERS.get(MIMES.get(mime, ""))
    if fn is None:
        raise ParseError(f"unsupported file type: {ext or mime or 'unknown'}")
    return fn(data, ocr) if fn is _pdf else fn(data)
